import io
import os
import cv2
import numpy as np
import tkinter as tk
import customtkinter as ctk
from PIL import Image, ImageTk
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.utils import get_column_letter
from tkinter import filedialog, messagebox

class Step2Frame(ctk.CTkFrame):
    def __init__(self, parent, app):
        super().__init__(parent, fg_color="transparent")
        self.app = app
        self._thumbs = []  # GC 방지용 리스트
        self._build_ui()

    def _build_ui(self):
        self.grid_rowconfigure(1, weight=1)
        self.grid_columnconfigure(0, weight=1)

        # Header bar
        hdr = ctk.CTkFrame(self, fg_color=("#e0e0e0", "#1a1d24"), corner_radius=0)
        hdr.grid(row=0, column=0, sticky="ew")

        ctk.CTkButton(hdr, text=self.app.app.t("ai_step2_back", "◀ 뒤로"), width=80,
                      fg_color=("#d0d0ff", "#2a2d44"), text_color=("black", "white"),
                      command=self.app.show_step1).pack(side="left", padx=15, pady=12)

        title_frame = ctk.CTkFrame(hdr, fg_color="transparent")
        title_frame.pack(side="left", padx=10)
        ctk.CTkLabel(title_frame, text=self.app.app.t("ai_step2_title", "크롭 결과 미리보기"), font=ctk.CTkFont("Segoe UI", 18, "bold")).pack(anchor="w")
        ctk.CTkLabel(title_frame, text=self.app.app.t("ai_step2_subtitle", "설정된 정렬 방식대로 엑셀에 저장됩니다"), text_color="gray").pack(anchor="w")

        right_frame = ctk.CTkFrame(hdr, fg_color="transparent")
        right_frame.pack(side="right", padx=15, pady=12)

        ctk.CTkLabel(right_frame, text=self.app.app.t("ai_step2_per_row", "한 열 개수 (세로):")).pack(side="left", padx=(0, 5))
        self.per_row_var = ctk.StringVar(value="6")
        ctk.CTkEntry(right_frame, textvariable=self.per_row_var, width=50, justify="center").pack(side="left", padx=(0, 15))
        self.per_row_var.trace_add("write", lambda *_: self.render_preview())

        ctk.CTkButton(right_frame, text=self.app.app.t("ai_step2_export", "📊 엑셀로 내보내기"), fg_color="#10b981", hover_color="#059669",
                      command=self._export_excel).pack(side="left")

        # Dual-axis Scrollable Area (Horizontal + Vertical)
        scroll_container = ctk.CTkFrame(self, fg_color="transparent")
        scroll_container.grid(row=1, column=0, sticky="nsew", padx=20, pady=20)
        scroll_container.grid_rowconfigure(0, weight=1)
        scroll_container.grid_columnconfigure(0, weight=1)

        self.canvas = tk.Canvas(scroll_container, bg=("#ebebeb" if ctk.get_appearance_mode()=="Light" else "#242424"), 
                                highlightthickness=0, borderwidth=0)
        self.canvas.grid(row=0, column=0, sticky="nsew")

        self.vsb = ctk.CTkScrollbar(scroll_container, orientation="vertical", command=self.canvas.yview)
        self.vsb.grid(row=0, column=1, sticky="ns")
        self.hsb = ctk.CTkScrollbar(scroll_container, orientation="horizontal", command=self.canvas.xview)
        self.hsb.grid(row=1, column=0, sticky="ew")

        self.canvas.configure(yscrollcommand=self.vsb.set, xscrollcommand=self.hsb.set)
        
        self.scroll = ctk.CTkFrame(self.canvas, fg_color="transparent")
        self.canvas_window = self.canvas.create_window((0, 0), window=self.scroll, anchor="nw")
        
        self.scroll.bind("<Configure>", lambda e: self.canvas.configure(scrollregion=self.canvas.bbox("all")))
        self.canvas.bind("<Configure>", lambda e: self.canvas.itemconfig(self.canvas_window, height=max(e.height, self.scroll.winfo_reqheight())))

    def render_preview(self):
        for w in self.scroll.winfo_children(): w.destroy()
        self._thumbs.clear()
        try:
            per_row = max(1, int(self.per_row_var.get()))
        except: per_row = 6

        out_w, out_h = self.app.step1.get_output_size()
        
        # Auto-fit thumbnail calculation
        screen_h = self.winfo_screenheight()
        available_h = screen_h - 180
        thumb_bound_h = max(20, int(available_h / per_row - 45))
        thumb_bound_w = 400

        for i, f in enumerate(self.app.files):
            col_idx = i // per_row
            row_idx = (i % per_row) + 1 # row 0 is for column header
            
            # Add Column Header (#1, #2...) at row 0
            if i % per_row == 0:
                ctk.CTkLabel(self.scroll, text=f"#{col_idx + 1}", 
                             font=ctk.CTkFont(size=16, weight="bold"),
                             text_color="#60a5fa").grid(row=0, column=col_idx, pady=(10, 15))

            try:
                img = f["image"]
                if f.get("angle", 0) != 0: img = img.rotate(f["angle"], resample=Image.BICUBIC, expand=True)
                
                # AI Matching or fallback
                coords = self._get_crop_coords_with_ai(i, img)
                if not coords: continue
                
                cropped = img.crop(coords)
                
                # Image Preview
                thumb = cropped.copy()
                thumb.thumbnail((thumb_bound_w, thumb_bound_h), Image.LANCZOS)
                tk_img = ImageTk.PhotoImage(thumb)
                self._thumbs.append(tk_img)
                
                card = ctk.CTkFrame(self.scroll, fg_color=("#e8e8e8", "#1e2028"), corner_radius=8, cursor="hand2")
                card.grid(row=row_idx, column=col_idx, padx=6, pady=6, sticky="n")
                card.bind("<Button-1>", lambda e, ii=i: self._open_fix_dialog(ii))

                lbl_img = ctk.CTkLabel(card, text="", image=tk_img)
                lbl_img.pack(padx=8, pady=(8, 4))
                lbl_img.bind("<Button-1>", lambda e, ii=i: self._open_fix_dialog(ii))

                ctk.CTkLabel(card, text=f"#{i+1}  {f['name']}", font=ctk.CTkFont(size=10), text_color="gray").pack(padx=8, pady=(0, 8))
            except Exception as e:
                print(f"Preview error for {f['name']}: {e}")
        
        # Reset scroll position
        self.canvas.xview_moveto(0)
        self.canvas.yview_moveto(0)

    def _get_crop_coords_with_ai(self, idx, img_rotated):
        f = self.app.files[idx]
        # 1. Manual/AI pre-calculated rect
        if f.get("crop_rect"):
            x, y, w, h = f["crop_rect"]
            dw, dh = img_rotated.size
            return (max(0, int(x)), max(0, int(y)), min(dw, int(x+w)), min(dh, int(y+h)))
        
        # 2. Global relative fallback
        if self.app.center_rel_rect:
            rx, ry, rw, rh = self.app.center_rel_rect
            dw, dh = img_rotated.size
            x1 = max(0, dw/2 + rx*dw - (rw*dw/2))
            y1 = max(0, dh/2 + ry*dh - (rh*dw/2))
            return (int(x1), int(y1), int(min(dw, x1+rw*dw)), int(min(dh, y1+rh*dh)))
        return None

    def _export_excel(self):
        if not self.app.files: return
        path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel files", "*.xlsx")], initialfile="Batch_Result.xlsx")
        if not path: return

        try:
            per_row = int(self.per_row_var.get())
            out_w, out_h = self.app.step1.get_output_size()
            wb = Workbook()
            ws = wb.active
            ws.title = "Photos"

            # Excel measurement conversion
            row_height_pt = out_h * 0.75
            col_width_units = max(20, int(out_w / 7) + 2)

            num_cols = (len(self.app.files) + per_row - 1) // per_row
            for c in range(1, num_cols + 1):
                col_letter = get_column_letter(c)
                ws.column_dimensions[col_letter].width = col_width_units
                # [NEW] Group each column
                ws.column_dimensions[col_letter].outline_level = 1
                # Set Title at Row 1
                ws.cell(row=1, column=c, value=f"#{c}")

            for i, f in enumerate(self.app.files):
                img = f["image"]
                if f.get("angle", 0) != 0: img = img.rotate(f["angle"], resample=Image.BICUBIC, expand=True)
                
                coords = self._get_crop_coords_with_ai(i, img)
                if not coords: continue
                
                resized = img.crop(coords).resize((int(out_w), int(out_h)), Image.LANCZOS)
                buf = io.BytesIO()
                resized.save(buf, format="PNG")
                buf.seek(0)

                xl_img = XlImage(buf)
                xl_img.width, xl_img.height = out_w, out_h
                
                col_idx = i // per_row
                row_idx = i % per_row
                
                # Each image takes 1 Excel row. Row 1 is Title.
                c_row = row_idx + 2 
                c_col = get_column_letter(col_idx + 1)
                
                # Sync Excel row height to image height to remove gaps
                ws.row_dimensions[c_row].height = row_height_pt
                
                ws.add_image(xl_img, f"{c_col}{c_row}")

            wb.save(path)
            self._show_success_dialog(path)
        except Exception as e:
            messagebox.showerror(self.app.app.t("ai_step2_error_title", "수행 오류"), f"{self.app.app.t('ai_step2_error_msg', '엑셀 저장 중 오류가 발생했습니다:')} {e}")

    def _show_success_dialog(self, path):
        dialog = ctk.CTkToplevel(self); dialog.title(self.app.app.t("ai_step2_success_title", "완료")); dialog.geometry("300x120"); dialog.grab_set()
        ctk.CTkLabel(dialog, text=self.app.app.t("ai_step2_success_msg", "✅ 엑셀 파일이 저장되었습니다!"), font=ctk.CTkFont(size=14, weight="bold")).pack(pady=(20, 5))
        ctk.CTkLabel(dialog, text=path, text_color="gray", font=ctk.CTkFont(size=10), wraplength=260).pack()
        ctk.CTkButton(dialog, text="OK", width=80, command=dialog.destroy).pack(pady=10)

    def _open_fix_dialog(self, idx):
        f = self.app.files[idx]
        dialog = ctk.CTkToplevel(self); dialog.title(f"{self.app.app.t('ai_step2_fix_title', '확대 보기 및 보정')} - {f['name']}"); dialog.geometry("700x850"); dialog.grab_set()
        
        curr_angle = tk.DoubleVar(value=f.get("angle", 0))
        subtitle = self.app.app.t("ai_step2_fix_subtitle", "사진 #{idx} 상세 보정").format(idx=idx+1)
        ctk.CTkLabel(dialog, text=subtitle, font=ctk.CTkFont(size=18, weight="bold")).pack(pady=15)
        
        canvas_size = 600
        canvas = tk.Canvas(dialog, bg="#111111", width=canvas_size, height=canvas_size, highlightthickness=0)
        canvas.pack(padx=20, pady=5)
        
        def update_view(val=None):
            img = f["image"].rotate(curr_angle.get(), resample=Image.BICUBIC, expand=True)
            c = self._get_crop_coords_with_ai(idx, img)
            if c:
                cropped = img.crop(c)
                iw, ih = cropped.size
                scale = min(canvas_size/iw, canvas_size/ih)
                disp = cropped.resize((int(iw*scale), int(ih*scale)), Image.LANCZOS)
                photo = ImageTk.PhotoImage(disp)
                canvas.image = photo
                canvas.delete("all")
                canvas.create_image(canvas_size//2, canvas_size//2, anchor="center", image=photo)
        
        update_view()
        ctk.CTkSlider(dialog, from_=-15, to=15, variable=curr_angle, command=update_view).pack(fill="x", padx=60, pady=15)
        ctk.CTkLabel(dialog, text=self.app.app.t("ai_step2_fix_slider", "슬라이더를 움직여 수평을 미세 조정하세요")).pack()
        
        def save_close():
            f["angle"] = curr_angle.get(); dialog.destroy(); self.render_preview()
        ctk.CTkButton(dialog, text=self.app.app.t("ai_step2_fix_done", "보정 완료"), fg_color="#10b981", command=save_close).pack(pady=25)
