from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import datetime

C_BG       = RGBColor(0xF5, 0xF7, 0xFA)
C_HEADER   = RGBColor(0x1A, 0x3C, 0x6E)
C_ACCENT   = RGBColor(0x00, 0x7A, 0xC3)
C_OK       = RGBColor(0x21, 0x96, 0x53)
C_WARN     = RGBColor(0xE6, 0x7E, 0x22)
C_WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
C_DARK     = RGBColor(0x1A, 0x1A, 0x2E)
C_LIGHT_BG = RGBColor(0xE8, 0xF0, 0xFB)
C_BORDER   = RGBColor(0xC9, 0xD4, 0xE8)
C_AI_BG    = RGBColor(0xEA, 0xF4, 0xFF)
C_AI_LINE  = RGBColor(0x00, 0x7A, 0xC3)

def add_rect(slide, x, y, w, h, fill_rgb=None, line_rgb=None, line_width=Pt(0.5)):
    shape = slide.shapes.add_shape(1, x, y, w, h)
    if fill_rgb:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_rgb
    else:
        shape.fill.background()
    if line_rgb:
        shape.line.color.rgb = line_rgb
        shape.line.width = line_width
    else:
        shape.line.fill.background()
    return shape

def add_text_box(slide, text, x, y, w, h, font_size=Pt(11), bold=False,
                 color=C_DARK, align=PP_ALIGN.LEFT, italic=False):
    txBox = slide.shapes.add_textbox(x, y, w, h)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = font_size
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.italic = italic
    return txBox

def make_card(slide, x, y, w, h, title, items,
              title_color=C_ACCENT, bg_color=C_LIGHT_BG, item_size=Pt(10)):
    add_rect(slide, x, y, w, h, fill_rgb=bg_color, line_rgb=C_BORDER)
    add_text_box(slide, title,
                 x + Inches(0.1), y + Inches(0.05),
                 w - Inches(0.2), Inches(0.28),
                 font_size=Pt(10), bold=True, color=title_color)
    txBox = slide.shapes.add_textbox(
        x + Inches(0.12), y + Inches(0.32),
        w - Inches(0.22), h - Inches(0.38))
    tf = txBox.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
        p.space_before = Pt(1.5)
        run = p.add_run()
        run.text = item
        run.font.size = item_size
        run.font.color.rgb = C_DARK

# ====================================================================
prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)

slide = prs.slides.add_slide(prs.slide_layouts[6])

# 배경
add_rect(slide, 0, 0, prs.slide_width, prs.slide_height, fill_rgb=C_BG)

# ── 헤더 ─────────────────────────────────────────────────────────
add_rect(slide, 0, 0, prs.slide_width, Inches(1.3), fill_rgb=C_HEADER)

add_text_box(slide,
             "AI 툴 활용 계측기 자동화 개발 보고",
             Inches(0.35), Inches(0.08),
             Inches(10.5), Inches(0.58),
             font_size=Pt(26), bold=True, color=C_WHITE)

add_text_box(slide,
             "Claude AI를 활용, 9종 계측기 연동·측정·기록 자동화 웹 앱 구축",
             Inches(0.38), Inches(0.68),
             Inches(10.5), Inches(0.38),
             font_size=Pt(12), color=RGBColor(0xB8, 0xD0, 0xF0))

today = datetime.date.today().strftime("%Y. %m. %d")
add_text_box(slide, today,
             Inches(11.0), Inches(0.10),
             Inches(2.1), Inches(0.35),
             font_size=Pt(11), color=RGBColor(0xB8, 0xD0, 0xF0),
             align=PP_ALIGN.RIGHT)

# ── AI 활용 방법 스트립 ──────────────────────────────────────────
add_rect(slide, 0, Inches(1.3), prs.slide_width, Inches(0.48),
         fill_rgb=RGBColor(0xDC, 0xE8, 0xF8))

add_text_box(slide,
             "▶  계측기별 통신 프로토콜(SCPI/Serial) 분석 · 코드 생성 · 오류 진단을 Claude AI와 대화 방식으로 진행 → 개발 기간 단축 및 반복 디버깅 자동화",
             Inches(0.3), Inches(1.33),
             Inches(12.7), Inches(0.4),
             font_size=Pt(10.5), color=C_HEADER)

# ── 본문 3열 ─────────────────────────────────────────────────────
TOP    = Inches(1.87)
CH     = Inches(3.95)
GAP    = Inches(0.16)
COL1W  = Inches(3.75)
COL2W  = Inches(4.2)
COL3W  = Inches(4.4)
X1 = Inches(0.18)
X2 = X1 + COL1W + GAP
X3 = X2 + COL2W + GAP

# 카드1: 지원 계측기
make_card(slide, X1, TOP, COL1W, CH,
          "◆  지원 계측기  ( 9 종 )",
          [
              "  저항계        Hioki 3540",
              "                Keithley 2700",
              "  절연저항      Agilent 4339B",
              "  데이터수집    DAQ-6510 (멀티채널)",
              "  거리계        Mitutoyo VL-50",
              "  박리·인장     SP-2100 / TL-2200",
              "  프로브 택     PT-2000",
              "  루프 택       LT-1000",
              "  이미지 편집   AI Photo Editor",
          ])

# 카드2: AI 활용 개발 내용
make_card(slide, X2, TOP, COL2W, CH,
          "◆  AI 툴 활용 내용 (Claude)",
          [
              "✔  계측기 통신 프로토콜 분석 및 코드 생성",
              "✔  SCPI 명령 오류(-113/-213) 원인 진단",
              "✔  측정 시퀀스 자동 검증 및 재작성",
              "✔  UI 레이아웃 · 그래프 · 테이블 구현",
              "✔  그룹별 통계 수염차트 알고리즘 생성",
              "✔  재측정 덮어쓰기 / 설정 영속화 구현",
              "✔  EXE 단독 배포 빌드 파이프라인 구축",
              "✔  코드 리뷰 · 버그 수정 · 문서화 자동화",
          ])

# 카드3: 개발 현황 (컬러 직접 구성)
add_rect(slide, X3, TOP, COL3W, CH,
         fill_rgb=C_LIGHT_BG, line_rgb=C_BORDER)
add_text_box(slide, "◆  개발 현황",
             X3 + Inches(0.1), TOP + Inches(0.05),
             COL3W - Inches(0.2), Inches(0.28),
             font_size=Pt(10), bold=True, color=C_ACCENT)

txBox3 = slide.shapes.add_textbox(
    X3 + Inches(0.12), TOP + Inches(0.35),
    COL3W - Inches(0.22), CH - Inches(0.42))
tf3 = txBox3.text_frame
tf3.word_wrap = True

def row(tf, text, size=Pt(10), bold=False, color=C_DARK, before=Pt(2), first=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_before = before
    run = p.add_run()
    run.text = text
    run.font.size = size
    run.font.bold = bold
    run.font.color.rgb = color

row(tf3, "✅  실기 검증 완료", Pt(10), True, C_OK, Pt(0), first=True)
for item in ["Hioki 3540", "Keithley 2700", "Mitutoyo VL-50",
             "SP-2100 / TL-2200", "PT-2000", "LT-1000", "AI Photo Editor"]:
    row(tf3, "     • " + item)

row(tf3, "🔧  실기 테스트 진행 예정", Pt(10), True, C_WARN, Pt(7))
row(tf3, "     • DAQ-6510 — 멀티채널 스캔 최종 확인")
row(tf3, "     • Agilent 4339B — 측정 시퀀스 최종 확인")

# ── AI 뱃지 (우하단 카드 안) ────────────────────────────────────
badge_x = X3 + Inches(0.12)
badge_y = TOP + CH - Inches(0.72)
badge_w = COL3W - Inches(0.24)
badge_h = Inches(0.58)
add_rect(slide, badge_x, badge_y, badge_w, badge_h,
         fill_rgb=RGBColor(0xE0, 0xF0, 0xFF), line_rgb=C_AI_LINE, line_width=Pt(1))
add_text_box(slide,
             "🤖  전체 개발 기간 내 Claude AI 지속 활용\n    코드 생성 · 디버깅 · 문서화 일괄 처리",
             badge_x + Inches(0.08), badge_y + Inches(0.06),
             badge_w - Inches(0.12), badge_h - Inches(0.08),
             font_size=Pt(9.5), color=C_ACCENT, bold=False)

# ── 하단 푸터 ────────────────────────────────────────────────────
add_rect(slide, 0, Inches(6.88), prs.slide_width, Inches(0.62),
         fill_rgb=C_HEADER)
add_text_box(slide, "3M Korea — 계측기 통합 관리 시스템  |  기술팀",
             Inches(0.35), Inches(6.91), Inches(8), Inches(0.38),
             font_size=Pt(10), color=RGBColor(0xB8, 0xD0, 0xF0))
add_text_box(slide, "Confidential",
             Inches(11.3), Inches(6.91), Inches(1.8), Inches(0.38),
             font_size=Pt(10), italic=True,
             color=RGBColor(0xB8, 0xD0, 0xF0), align=PP_ALIGN.RIGHT)

# ====================================================================
out = r"c:\Users\0op64\.gemini\antigravity\scratch\3M Instrument Logger\3M_Instrument_Logger_보고.pptx"
prs.save(out)
print(f"저장 완료: {out}")
